import argparse
import os
from zipfile import BadZipFile
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_openai import OpenAIEmbeddings


DATA_DIR = Path(__file__).parent / "sample_prd"
CHUNK_MAX_CHARS = 420
CHUNK_OVERLAP_PARAGRAPHS = 1
TOP_K = 5
NEIGHBOR_WINDOW = 1


@dataclass
class Paragraph:
    text: str
    index: int
    section: str
    is_heading: bool


EVALUATION_CASES = [
    {
        "question": "合作伙伴入驻有哪些注意事项？",
        "expected_sources": {
            "星云服务平台合作伙伴入驻指南.docx"
        },
    },
    {
        "question": "为什么要做 AI 智能回复？",
        "expected_sources": {"星云服务平台AI助手需求.docx"},
    },
    {
        "question": "智能回复 Agent 需要具备哪些能力？",
        "expected_sources": {"星云服务平台AI助手需求.docx"},
    },
    {
        "question": "需求文档中是否规划了“小憩”状态？前后台分别如何设计？",
        "expected_sources": {
            "星云服务平台V2需求.docx",
            "星云服务平台V3需求.docx",
        },
    },
    {
        "question": "客户小程序如何入驻？",
        "expected_sources": set(),
    },
]


DIAGNOSTIC_KEYWORDS = {
    "问题 2：为什么要做 AI 智能回复？": [
        "渠道工程师专业能力有高有低",
        "自主闭环大多数P3P4类的问题",
        "自动生成专业建议",
    ],
    "问题 3：智能回复 Agent 需要具备哪些能力？": [
        "确认诉求",
        "开放式结尾",
        "风险提示",
        "核心步骤拆解",
        "验证方法",
    ],
}


def is_heading(style_name: str) -> bool:
    normalized = style_name.lower()
    return "heading" in normalized or "标题" in style_name


def read_paragraphs(path: Path) -> list[Paragraph]:
    """读取 DOCX 正文，并记录段落序号及最近的标题。"""
    docx = DocxDocument(path)
    paragraphs: list[Paragraph] = []
    current_section = "文档开头"

    for index, paragraph in enumerate(docx.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""
        paragraph_is_heading = is_heading(style_name)
        if paragraph_is_heading:
            current_section = text
            # 标题会写入 chunk 前缀和 metadata，不作为独立正文段落入库。
            continue

        paragraphs.append(
            Paragraph(
                text=text,
                index=index,
                section=current_section,
                is_heading=False,
            )
        )

    return paragraphs


def make_chunk(path: Path, paragraphs: list[Paragraph]) -> Document:
    first = paragraphs[0]
    last = paragraphs[-1]
    body = "\n".join(paragraph.text for paragraph in paragraphs)
    # 标题也进入 Embedding 文本，避免只存 metadata 导致检索时不可见。
    content = (
        f"[文档：{path.stem}]\n"
        f"[章节：{first.section}]\n"
        f"{body}"
    )

    return Document(
        page_content=content,
        metadata={
            "source": path.name,
            "section": first.section,
            "paragraph_start": first.index,
            "paragraph_end": last.index,
        },
    )


def chunk_paragraphs(
    path: Path,
    paragraphs: list[Paragraph],
    max_chars: int = CHUNK_MAX_CHARS,
    overlap_paragraphs: int = CHUNK_OVERLAP_PARAGRAPHS,
) -> list[Document]:
    """按段落组合 chunk，避免在段落中间机械截断。"""
    chunks: list[Document] = []
    buffer: list[Paragraph] = []
    buffer_chars = 0

    for paragraph in paragraphs:
        paragraph_chars = len(paragraph.text)

        section_changed = (
            buffer
            and paragraph.section != buffer[-1].section
        )
        exceeds_limit = (
            buffer
            and buffer_chars + paragraph_chars > max_chars
        )

        if section_changed or exceeds_limit:
            chunks.append(make_chunk(path, buffer))
            # 章节切换时不跨章节重叠；长度切分时保留少量段落上下文。
            if section_changed:
                buffer = []
            else:
                buffer = (
                    buffer[-overlap_paragraphs:]
                    if overlap_paragraphs
                    else []
                )
            buffer_chars = sum(len(item.text) for item in buffer)

        buffer.append(paragraph)
        buffer_chars += paragraph_chars

    if buffer:
        chunks.append(make_chunk(path, buffer))

    return chunks


def load_chunks() -> list[Document]:
    chunks: list[Document] = []

    for path in sorted(DATA_DIR.glob("*.docx")):
        if path.name.startswith("~$"):
            continue

        try:
            paragraphs = read_paragraphs(path)
        except (BadZipFile, OSError) as error:
            print(f"跳过无法读取的文档 {path.name}：{error}")
            continue

        document_chunks = chunk_paragraphs(path, paragraphs)
        for chunk_index, chunk in enumerate(document_chunks):
            chunk.metadata["chunk_index"] = chunk_index
        chunks.extend(document_chunks)
        print(
            f"{path.name}: "
            f"{len(paragraphs)} 个非空段落，"
            f"{len(document_chunks)} 个 chunk"
        )

    return chunks


def expand_with_neighbors(
    retrieved_documents: list[Document],
    all_chunks: list[Document],
    window: int = NEIGHBOR_WINDOW,
) -> list[Document]:
    """为召回的小 chunk 补充同文档、同章节的相邻 chunk。"""
    chunk_lookup = {
        (
            chunk.metadata["source"],
            chunk.metadata["chunk_index"],
        ): chunk
        for chunk in all_chunks
    }
    expanded: list[Document] = []
    seen: set[tuple[str, int]] = set()

    for document in retrieved_documents:
        source = document.metadata["source"]
        section = document.metadata["section"]
        center_index = document.metadata["chunk_index"]

        for neighbor_index in range(
            center_index - window,
            center_index + window + 1,
        ):
            neighbor = chunk_lookup.get((source, neighbor_index))
            if neighbor is None:
                continue
            if neighbor.metadata["section"] != section:
                continue

            key = (source, neighbor_index)
            if key not in seen:
                expanded.append(neighbor)
                seen.add(key)

    return expanded


def print_result(rank: int, document: Document, score: float) -> None:
    metadata = document.metadata
    preview = document.page_content.replace("\n", " ")[:220]

    print(f"\n  Top {rank} | score={score:.4f}")
    print(
        "  来源："
        f"{metadata['source']} | "
        f"章节：{metadata['section']} | "
        f"chunk：{metadata['chunk_index']} | "
        f"段落：{metadata['paragraph_start']}-{metadata['paragraph_end']}"
    )
    print(f"  内容：{preview}")


def normalize_text(text: str) -> str:
    """忽略空白差异，便于定位被 Word 拆开的中文文本。"""
    return "".join(text.split()).lower()


def diagnose_chunks(chunks: list[Document]) -> None:
    """定位标准答案关键词是否进入 chunk，不调用模型或 Embedding API。"""
    for question, keywords in DIAGNOSTIC_KEYWORDS.items():
        print(f"\n{'=' * 70}")
        print(question)

        for keyword in keywords:
            normalized_keyword = normalize_text(keyword)
            matches = [
                chunk
                for chunk in chunks
                if normalized_keyword in normalize_text(chunk.page_content)
            ]

            print(f"\n关键词：{keyword}")
            print(f"命中 chunk 数：{len(matches)}")

            for chunk in matches:
                metadata = chunk.metadata
                preview = chunk.page_content.replace("\n", " ")[:500]
                print(
                    "  来源："
                    f"{metadata['source']} | "
                    f"章节：{metadata['section']} | "
                    f"段落：{metadata['paragraph_start']}-"
                    f"{metadata['paragraph_end']}"
                )
                print(f"  内容：{preview}")


def run_retrieval(chunks: list[Document]) -> None:
    embeddings = OpenAIEmbeddings(
        model="text-embedding-v4",
        api_key=os.getenv("DASHSCOPE_API_KEY"),
        base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
        check_embedding_ctx_length=False,
        # 当前百炼 Embedding 接口单批最多接受 10 条文本。
        chunk_size=10,
    )

    vector_store = InMemoryVectorStore.from_documents(
        documents=chunks,
        embedding=embeddings,
    )

    for index, case in enumerate(EVALUATION_CASES, start=1):
        question = case["question"]
        expected_sources = case["expected_sources"]
        results = vector_store.similarity_search_with_score(
            question,
            k=TOP_K,
        )

        print(f"\n{'=' * 70}")
        print(f"问题 {index}：{question}")

        for rank, (document, score) in enumerate(results, start=1):
            print_result(rank, document, score)

        expanded_documents = expand_with_neighbors(
            [document for document, _ in results],
            chunks,
        )
        print(
            f"\n  邻居扩展后：{len(expanded_documents)} 个 chunk"
        )
        for document in expanded_documents:
            metadata = document.metadata
            print(
                "  - "
                f"{metadata['source']} | "
                f"{metadata['section']} | "
                f"chunk {metadata['chunk_index']} | "
                f"段落 {metadata['paragraph_start']}-"
                f"{metadata['paragraph_end']}"
            )

        retrieved_sources = {
            document.metadata["source"]
            for document, _ in results
        }

        if expected_sources:
            matched_sources = expected_sources & retrieved_sources
            print(
                "\n  期望来源命中："
                f"{len(matched_sources)}/{len(expected_sources)} "
                f"{sorted(matched_sources)}"
            )
        else:
            print(
                "\n  这是无答案样本：即使向量库返回 top-k，"
                "也不能据此认定资料足以回答。"
            )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="需求文档检索基线与 chunk 诊断"
    )
    parser.add_argument(
        "--diagnose",
        action="store_true",
        help="只定位问题 2、3 的答案关键词，不调用 Embedding API",
    )
    args = parser.parse_args()

    chunks = load_chunks()
    print(f"\n总计：{len(chunks)} 个 chunk")

    if args.diagnose:
        diagnose_chunks(chunks)
        return

    run_retrieval(chunks)


if __name__ == "__main__":
    main()
